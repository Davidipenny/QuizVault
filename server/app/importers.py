from __future__ import annotations

import base64
import io
import importlib.util
import json
import re
from pathlib import Path
from typing import Any

from .domain import compact_text


TYPE_ALIASES = {
    "单选": "single", "单选题": "single", "single": "single",
    "多选": "multi", "多选题": "multi", "multi": "multi",
    "任意选": "any", "不定项": "any", "any": "any",
    "判断": "truefalse", "判断题": "truefalse", "truefalse": "truefalse",
    "填空": "fill", "填空题": "fill", "fill": "fill",
    "问答": "essay", "问答题": "essay", "简答": "essay", "essay": "essay",
}


def normalize_type(value: Any) -> str:
    return TYPE_ALIASES.get(compact_text(str(value or "")).lower(), str(value or "").lower())


def row_to_question(row: dict[str, Any], source: str = "") -> dict:
    qtype = normalize_type(row.get("type") or row.get("题型"))
    prompt = row.get("prompt") or row.get("question") or row.get("题干") or ""
    explanation = row.get("explanation") or row.get("解析") or ""
    case_material = row.get("case_material") or row.get("材料") or ""
    choices = row.get("choices") or []
    if isinstance(choices, dict):
        choices = [{"label": k, "content": v} for k, v in choices.items()]
    if not choices:
        for label in "ABCDEFGHIJ":
            value = row.get(label) or row.get(f"选项{label}")
            if value not in (None, ""):
                choices.append({"label": label, "content": str(value)})
    answer = row.get("answer_spec") or row.get("answer") or row.get("答案")
    if not isinstance(answer, dict):
        text = compact_text(str(answer or ""))
        if qtype in {"single", "multi", "any"}:
            answer_spec = {"correct": re.findall(r"[A-J]", text.upper())}
        elif qtype == "truefalse":
            answer_spec = {"value": text.lower() in {"a", "true", "正确", "对", "1", "是"}}
        elif qtype == "fill":
            blanks = [[part.strip()] for part in re.split(r"[|；;]", text) if part.strip()]
            answer_spec = {"blanks": blanks, "unordered": bool(row.get("unordered") or row.get("答案无序"))}
        else:
            answer_spec = {"reference": text}
    else:
        answer_spec = answer
    correct = set(answer_spec.get("correct", []))
    for index, choice in enumerate(choices):
        choice.setdefault("label", chr(65 + index))
        choice["label"] = str(choice["label"]).upper()
        choice["content"] = str(choice.get("content", ""))
        choice["is_correct"] = choice["label"] in correct
        choice["display_order"] = index
    return {
        "type": qtype,
        "prompt": str(prompt),
        "case_material": str(case_material),
        "explanation": str(explanation),
        "source": source,
        "answer_spec": answer_spec,
        "choices": choices,
    }


def parse_json(content: str) -> list[dict]:
    data = json.loads(content)
    if isinstance(data, dict):
        data = data.get("questions", [data])
    if not isinstance(data, list):
        raise ValueError("JSON 顶层必须是题目数组或包含 questions 数组的对象")
    return [row_to_question(row, "AI JSON") for row in data]


def parse_delimited_text(content: str) -> list[dict]:
    stripped = content.strip()
    if stripped.startswith("[") or stripped.startswith("{"):
        return parse_json(stripped)
    blocks = [b.strip() for b in re.split(r"\n\s*---+\s*\n", stripped) if b.strip()]
    rows = []
    for block in blocks:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        row: dict[str, Any] = {}
        options: dict[str, str] = {}
        question_lines = []
        for line in lines:
            match = re.match(r"^(题型|类型|答案|解析|材料)\s*[:：]\s*(.*)$", line, re.I)
            option = re.match(r"^([A-J])[.、．:]\s*(.*)$", line, re.I)
            numbered = re.match(r"^\d+[.、．]\s*(.*)$", line)
            if match:
                key, value = match.groups()
                mapping = {"题型": "type", "类型": "type", "答案": "answer", "解析": "explanation", "材料": "case_material"}
                row[mapping[key]] = value
            elif option:
                options[option.group(1).upper()] = option.group(2)
            elif numbered and not question_lines:
                question_lines.append(numbered.group(1))
            else:
                question_lines.append(line)
        row["question"] = "\n".join(question_lines)
        row["choices"] = options
        rows.append(row_to_question(row, "文本导入"))
    return rows


def parse_excel(encoded: str, filename: str) -> list[dict]:
    raw = base64.b64decode(encoded)
    if filename.lower().endswith(".xls"):
        import xlrd
        book = xlrd.open_workbook(file_contents=raw)
        sheet = book.sheet_by_index(0)
        headers = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
        rows = [dict(zip(headers, [sheet.cell_value(r, c) for c in range(sheet.ncols)])) for r in range(1, sheet.nrows)]
    else:
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        values = list(sheet.iter_rows(values_only=True))
        headers = [str(x or "").strip() for x in values[0]] if values else []
        rows = [dict(zip(headers, row)) for row in values[1:]]
    return [row_to_question(row, filename) for row in rows if any(v not in (None, "") for v in row.values())]


def parse_word(encoded: str, filename: str) -> list[dict]:
    from docx import Document
    document = Document(io.BytesIO(base64.b64decode(encoded)))
    text = "\n".join(p.text for p in document.paragraphs)
    return parse_delimited_text(text)


def parse_source(source_type: str, content: str, filename: str = "") -> list[dict]:
    if source_type in {"json", "ai_json"}:
        return parse_json(content)
    if source_type == "text":
        return parse_delimited_text(content)
    if source_type == "excel":
        return parse_excel(content, filename)
    if source_type == "word":
        return parse_word(content, filename)
    raise ValueError(f"不支持的导入类型：{source_type}")


def load_legacy_bank(bank_dir: Path) -> list[dict]:
    parser_path = Path(__file__).resolve().parents[2] / "parse_markdown.py"
    spec = importlib.util.spec_from_file_location("quizvault_legacy_parser", parser_path)
    if not spec or not spec.loader:
        raise RuntimeError("无法加载旧版 Markdown 解析器")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    parse_markdown = module.parse_markdown

    result: list[dict] = []
    for file in sorted(bank_dir.iterdir()):
        if file.suffix.lower() == ".md":
            for item in parse_markdown(file.read_text(encoding="utf-8")):
                answer = str(item.get("answer", "")).upper()
                qtype = item.get("type", "single")
                choices = item.get("options", {})
                converted = row_to_question({
                    "type": qtype,
                    "question": item.get("question", ""),
                    "choices": choices,
                    "answer": answer,
                    "explanation": item.get("explanation", ""),
                }, file.name)
                converted["_legacy_id"] = item.get("id")
                result.append(converted)
    return result
